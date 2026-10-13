#!/usr/bin/env python3
"""LIBROAD 테스트용 문서 만들기 (Linux).  usage: make_samples.py <fake-home> [pyhwp fixtures dir]
python-docx, openpyxl, python-pptx, reportlab, Pillow, LibreOffice(soffice) 필요."""
import os, sys, shutil, subprocess, zipfile, struct, time
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont

home = sys.argv[1]
fixtures = sys.argv[2] if len(sys.argv) > 2 else ""
D = {k: os.path.join(home, k) for k in ("Documents", "Desktop", "Downloads")}
for d in D.values():
    os.makedirs(d, exist_ok=True)
docs = D["Documents"]
os.makedirs(os.path.join(docs, "회사", "2024 사업"), exist_ok=True)
os.makedirs(os.path.join(docs, "개인"), exist_ok=True)
os.makedirs(os.path.join(docs, "빈 폴더", "안쪽도 빈 폴더"), exist_ok=True)
FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"

def docx(path, paras):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "표 안의 글자"
    t.cell(0, 1).text = "셀두번째"
    d.save(path)

docx(os.path.join(docs, "회사", "2024 사업", "사업계획서_최종.docx"), ["2024년 사업계획서", "올해의 목표는 도서관 검색 서비스를 만드는 것입니다.", "Budget planning for LIBROAD project."])
docx(os.path.join(docs, "회사", "2024 사업", "사업계획서_최종2.docx"), ["2024년 사업계획서 (수정)", "올해의 목표는 도서관 검색 서비스를 만드는 것입니다. 일정 추가."])
docx(os.path.join(docs, "회사", "2024 사업", "사업계획서_수정본(1).docx"), ["사업계획서 수정본", "검토 의견 반영"])
docx(os.path.join(docs, "개인", "자기소개서.docx"), ["저는 꼼꼼한 사람입니다.", "취미는 독서와 산책입니다."])

wb = Workbook(); ws = wb.active; ws.title = "예산"
ws.append(["항목", "금액", "비고"]); ws.append(["사무용품", 120000, "볼펜과 노트"]); ws.append(["출장비", 350000, "부산 출장"])
ws2 = wb.create_sheet("인원"); ws2.append(["이름", "부서"]); ws2.append(["김도서", "기획팀"])
wb.save(os.path.join(docs, "회사", "예산표.xlsx"))

pr = Presentation(); s = pr.slides.add_slide(pr.slide_layouts[1])
s.shapes.title.text = "분기 실적 발표"; s.placeholders[1].text = "매출이 20% 늘었습니다\n신규 고객 확보"
s.notes_slide.notes_text_frame.text = "발표자 메모: 천천히 말하기"
pr.save(os.path.join(docs, "회사", "분기실적발표.pptx"))

open(os.path.join(docs, "메모.txt"), "w", encoding="utf-8").write("장보기 목록\n우유, 사과, 식빵\n도서관 반납일 금요일\n")
open(os.path.join(docs, "옛날메모.txt"), "wb").write("예전 컴퓨터에서 만든 한글 메모입니다. 비밀번호 힌트 아님".encode("cp949"))
open(os.path.join(docs, "README.md"), "w", encoding="utf-8").write("# 프로젝트 노트\n\n마크다운 문서도 찾을 수 있어요. **markdown** search\n")
open(os.path.join(docs, "회사", "거래처.csv"), "wb").write("회사,담당자,전화\n한빛상사,이영희,02-123-4567\n".encode("cp949"))

# 예전 형식 (doc, xls, ppt, rtf) — LibreOffice로 변환
conv = os.path.join(home, "_conv"); os.makedirs(conv, exist_ok=True)
docx(os.path.join(conv, "회의록_2023.docx"), ["2023년 회의록", "참석자: 김철수, 박영희", "안건: 예산 집행 점검", "결정 사항은 다음 회의에서 확정"])
wb = Workbook(); ws = wb.active; ws.title = "재고"; ws.append(["품목", "수량"]); ws.append(["모니터", 15]); ws.append(["키보드 세트", 30]); wb.save(os.path.join(conv, "재고현황.xlsx"))
pr = Presentation(); s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = "옛날 발표 자료"; s.placeholders[1].text = "파워포인트 97 형식 시험"; pr.save(os.path.join(conv, "옛날발표.pptx"))
docx(os.path.join(conv, "안내문.docx"), ["리치 텍스트 안내문", "RTF format 문서입니다."])
def soffice(fmt, src, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", fmt, "--outdir", outdir, src], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
soffice("doc", os.path.join(conv, "회의록_2023.docx"), os.path.join(docs, "회사"))
soffice("xls", os.path.join(conv, "재고현황.xlsx"), os.path.join(docs, "회사"))
soffice("ppt", os.path.join(conv, "옛날발표.pptx"), os.path.join(docs, "회사"))
soffice("rtf", os.path.join(conv, "안내문.docx"), docs)
# 글자가 있는 PDF
docx(os.path.join(conv, "연구보고서.docx"), ["연구 보고서", "인공지능 기반 문서 분류 연구의 결과를 정리합니다.", "결론: 초성 검색이 편리하다."] + ["본문 문단 %d번입니다. 도서관 정리 이야기." % i for i in range(40)])
soffice("pdf", os.path.join(conv, "연구보고서.docx"), docs)
shutil.rmtree(conv)

# 스캔 문서 (글자가 그림으로만 있는 PDF) · 사진
def text_image(lines, size=(1240, 1754)):
    img = Image.new("RGB", size, "white"); dr = ImageDraw.Draw(img)
    f = ImageFont.truetype(FONT, 44)
    y = 120
    for l in lines:
        dr.text((110, y), l, fill="black", font=f); y += 80
    return img
text_image(["임대차 계약서", "임대인 홍길동 임차인 김영수", "보증금 일천만원정", "계약 기간 2년"]).save(os.path.join(D["Downloads"], "스캔_계약서.pdf"), "PDF", resolution=150)
text_image(["영수증", "카페 도서관점", "아메리카노 4,500원"], (900, 700)).save(os.path.join(D["Desktop"], "영수증사진.jpg"), quality=90)
text_image(["Screenshot test", "회의 일정 변경 안내"], (1000, 500)).save(os.path.join(D["Desktop"], "스크린샷 2024-05-02.png"))

# hwpx (한글 2014 이후 형식)
def hwpx(path, paras):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip")
        z.writestr("version.xml", '<?xml version="1.0"?><hv:HCFVersion xmlns:hv="http://www.hancom.co.kr/hwpml/2011/version" major="5"/>')
        body = "".join('<hp:p><hp:run><hp:t>%s</hp:t></hp:run></hp:p>' % p for p in paras)
        z.writestr("Contents/section0.xml", '<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">%s</hs:sec>' % body)
        z.writestr("Preview/PrvText.txt", "\n".join(paras))
hwpx(os.path.join(docs, "회사", "공문_협조요청.hwpx"), ["수신: 각 부서장", "제목: 자료 제출 협조 요청", "붙임 서류를 금요일까지 제출해 주시기 바랍니다."])
hwpx(os.path.join(D["Downloads"], "가정통신문.hwpx"), ["가정통신문", "현장 체험 학습 안내", "준비물: 도시락, 물"])

# hwp (한글 5.0) — pyhwp 테스트 문서
if fixtures and os.path.isdir(fixtures):
    for n in ["footnote-endnote.hwp", "headerfooter.hwp", "aligns.hwp", "issue30.hwp"]:
        p = os.path.join(fixtures, n)
        if os.path.exists(p):
            shutil.copy(p, os.path.join(docs, "회사", n))
    p = os.path.join(fixtures, "sample-5017.hwp")
    if os.path.exists(p):
        shutil.copy(p, os.path.join(docs, "회사", "샘플_5017.hwp"))

# epub
with zipfile.ZipFile(os.path.join(docs, "개인", "작은책.epub"), "w") as z:
    z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
    z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
    z.writestr("OEBPS/content.opf", '<?xml version="1.0"?><package xmlns:dc="http://purl.org/dc/elements/1.1/"><metadata><dc:title>작은 도서관 이야기</dc:title><dc:creator>무명 작가</dc:creator></metadata></package>')
    z.writestr("OEBPS/ch1.xhtml", '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>1장</title><style>p{}</style></head><body><h1>1장 서가 사이에서</h1><p>책장 사이로 햇살이 들었다.&nbsp;사서는 조용히 웃었다.</p></body></html>')

# iWork (Pages · Numbers · Keynote) — IWA 형식을 흉내 낸 작은 문서
def varint(n):
    out = b""
    while True:
        b = n & 0x7f; n >>= 7
        out += bytes([b | (0x80 if n else 0)])
        if not n: return out
def field(num, wire, payload):
    key = varint((num << 3) | wire)
    if wire == 2: return key + varint(len(payload)) + payload
    return key + varint(payload)
def snappy_literal(data):
    # 압축하지 않은 snappy 블록 (literal 하나)
    out = varint(len(data)); n = len(data) - 1
    if n < 60: out += bytes([n << 2])
    elif n < 256: out += bytes([60 << 2, n])
    else: out += bytes([61 << 2, n & 0xff, n >> 8])
    return out + data
def iwa(objects):
    stream = b""
    for typ, msg in objects:
        mi = field(1, 0, typ) + field(3, 0, len(msg))
        info = field(1, 0, 1) + field(2, 2, mi)
        stream += varint(len(info)) + info + msg
    comp = snappy_literal(stream)
    return b"\x00" + struct.pack("<I", len(comp))[:3] + comp
storage = field(1, 0, 0) + field(3, 2, "아이워크 페이지스 문서 본문입니다\n가을 독서 모임 계획".encode())
with zipfile.ZipFile(os.path.join(docs, "개인", "독서모임.pages"), "w") as z:
    z.writestr("Index/Document.iwa", iwa([(2001, storage), (2002, field(3, 2, b"noise-should-skip"))]))
tdl = field(1, 0, 1) + field(2, 0, 3) + field(3, 2, field(1, 0, 1) + field(2, 0, 1) + field(3, 2, "가계부 식비 항목".encode())) + field(3, 2, field(1, 0, 2) + field(2, 0, 1) + field(3, 2, "교통비".encode()))
with zipfile.ZipFile(os.path.join(docs, "개인", "가계부.numbers"), "w") as z:
    z.writestr("Index/Tables/DataList.iwa", iwa([(6005, tdl)]))
# Keynote 패키지 (폴더 형식)
kd = os.path.join(docs, "회사", "신제품 소개.key", "Index"); os.makedirs(kd, exist_ok=True)
open(os.path.join(kd, "Slide.iwa"), "wb").write(iwa([(2001, field(3, 2, "키노트 신제품 소개 슬라이드".encode()))]))

# 중복 · 큰 파일 · 오래된 파일 · iCloud에만 있는 파일
shutil.copy(os.path.join(docs, "메모.txt"), os.path.join(D["Desktop"], "메모 복사본.txt"))
shutil.copy(os.path.join(docs, "메모.txt"), os.path.join(D["Downloads"], "메모.txt"))
with open(os.path.join(D["Downloads"], "큰자료.pdf"), "wb") as f:
    f.write(b"%PDF-1.4\n%fake big\n"); f.write(os.urandom(60 * 1024 * 1024))
old = os.path.join(docs, "개인", "2015_가족여행.txt")
open(old, "w", encoding="utf-8").write("2015년 제주도 가족 여행 일정")
t = time.mktime((2015, 7, 1, 0, 0, 0, 0, 0, -1)); os.utime(old, (t, t))
open(os.path.join(docs, "개인", "클라우드문서.icloudtest.docx"), "wb").write(b"PK\x03\x04 not downloaded")
print("samples ready in", home)
